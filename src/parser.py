"""
Intake Parsers for Resume and Job Description inputs across formats (Text, Markdown, DOCX, PDF, URL).
"""
import os
import re
from pathlib import Path
from typing import Union
import requests
from bs4 import BeautifulSoup

try:
    import docx
except ImportError:
    docx = None

try:
    import pypdf
except ImportError:
    pypdf = None


class InputParser:
    """Handles loading and parsing of text, markdown, docx, pdf, and web URLs."""

    @staticmethod
    def parse_file(file_path: Union[str, Path]) -> str:
        """Parse content from a local file (.txt, .md, .docx, .pdf)."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = path.suffix.lower()

        if suffix in [".txt", ".md", ".markdown"]:
            return path.read_text(encoding="utf-8", errors="replace").strip()

        elif suffix == ".docx":
            if docx is None:
                raise ImportError("python-docx is required to parse .docx files. Run: pip install python-docx")
            doc = docx.Document(path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n\n".join(full_text).strip()

        elif suffix == ".pdf":
            if pypdf is None:
                raise ImportError("pypdf is required to parse .pdf files. Run: pip install pypdf")
            reader = pypdf.PdfReader(path)
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            return "\n\n".join(pages_text).strip()

        else:
            # Fallback to plain text read
            try:
                return path.read_text(encoding="utf-8", errors="replace").strip()
            except Exception as e:
                raise ValueError(f"Unsupported file format '{suffix}' for {file_path}: {e}")

    @staticmethod
    def parse_url(url: str, timeout: int = 15) -> str:
        """Fetch and extract clean readable text from a job posting URL."""
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
        resp = requests.get(url, headers=headers, timeout=timeout)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "html.parser")

        # Remove script, style, header, footer, nav tags
        for element in soup(["script", "style", "nav", "header", "footer", "noscript", "svg"]):
            element.decompose()

        # Prioritize job description containers if available
        candidates = soup.find_all(
            ["div", "section", "article"],
            class_=re.compile(r"(job[-_]?desc|posting|description|content|details|career)", re.I)
        )

        if candidates:
            # Pick longest container
            best_candidate = max(candidates, key=lambda c: len(c.get_text()))
            text = best_candidate.get_text(separator="\n", strip=True)
        else:
            text = soup.get_text(separator="\n", strip=True)

        # Collapse excess empty lines
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines)

    @classmethod
    def load_content(cls, input_str_or_path: str) -> str:
        """
        Universal content loader. Detects if input is a URL, existing file path, or raw text.
        """
        raw = input_str_or_path.strip()
        if raw.startswith("http://") or raw.startswith("https://"):
            return cls.parse_url(raw)

        # Check if local file
        if os.path.exists(raw):
            return cls.parse_file(raw)

        # Otherwise treat as raw input string
        return raw
